#!/usr/bin/env python3
"""
esporta-excel.py — genera il registro ordini in Excel dall'archivio richieste.

    python3 assistant/scripts/esporta-excel.py [file.xlsx]

Perche' un export e non una scrittura diretta a ogni ordine: se il file e' aperto
in Excel, un processo che prova a scriverlo fallisce o lo corrompe. L'archivio
JSONL e' il registro sicuro (una riga in coda, mai riscritto); il foglio Excel e'
una fotografia che si rigenera quando serve.
"""
import json, os, sys, datetime
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ARCHIVIO = os.environ.get('CC_ARCHIVIO', os.path.join(BASE, 'data', 'richieste.jsonl'))
USCITA = sys.argv[1] if len(sys.argv) > 1 else os.path.join(BASE, 'data', 'registro-ordini.xlsx')

SCURO = '1A1610'
ORO = 'C4922A'
CHIARO = 'F5EEE2'

LINEE = {
    'liturgico': 'Articoli Liturgici',
    'garden': 'Linea Garden',
    'home-collection': 'Home Collection',
    'private-label': 'Private Label',
    'collaborazioni': 'Collaborazioni',
    'campionario': 'Campionario',
    'altro': 'Altro',
}

def leggi():
    if not os.path.exists(ARCHIVIO):
        return []
    voci = []
    with open(ARCHIVIO, encoding='utf-8') as f:
        for riga in f:
            riga = riga.strip()
            if not riga:
                continue
            try:
                voci.append(json.loads(riga))
            except json.JSONDecodeError:
                continue          # una riga rotta non deve far saltare l'export
    return voci

def intesta(ws, colonne):
    for i, (titolo, larghezza) in enumerate(colonne, start=1):
        c = ws.cell(row=1, column=i, value=titolo)
        c.font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor=SCURO)
        c.alignment = Alignment(vertical='center')
        ws.column_dimensions[get_column_letter(i)].width = larghezza
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = 'A2'

def scrivi(ws, riga, valori, grassetto=False):
    for i, v in enumerate(valori, start=1):
        c = ws.cell(row=riga, column=i, value=v)
        c.font = Font(name='Arial', size=10, bold=grassetto)
        c.alignment = Alignment(vertical='top', wrap_text=(i in (9, 10)))

def data_breve(ts):
    try:
        return datetime.datetime.fromisoformat(ts.replace('Z', '+00:00')).strftime('%d/%m/%Y')
    except Exception:
        return ts[:10]

voci = leggi()
wb = Workbook()

# ── 1. Richieste ────────────────────────────────────────────────
ws = wb.active
ws.title = 'Richieste'
intesta(ws, [('Data', 12), ('Cliente', 26), ('Azienda / Ente', 26), ('Email', 30),
             ('Linea', 20), ('Cadenza', 30), ('Prossima', 12), ('Pezzi', 9), ('Articoli', 42)])

for r, v in enumerate(voci, start=2):
    art = v.get('articoli') or []
    pezzi = sum(int(a.get('quantita') or 0) for a in art)
    scrivi(ws, r, [
        data_breve(v.get('ts', '')),
        v.get('nome', ''),
        v.get('azienda', ''),
        v.get('email', ''),
        LINEE.get(v.get('linea', ''), v.get('linea', '')),
        (v.get('fornitura') or {}).get('descrizione', '').split('\n')[0],
        v.get('prossima', ''),
        pezzi if pezzi else None,
        ' · '.join('%s%s' % (a.get('nome', ''), ' ×%d' % a['quantita'] if a.get('quantita') else '')
                   for a in art),
    ])

ultima = len(voci) + 1
if voci:
    tot = ultima + 1
    ws.cell(row=tot, column=7, value='Totale').font = Font(name='Arial', size=10, bold=True)
    c = ws.cell(row=tot, column=8, value='=SUM(H2:H%d)' % ultima)
    c.font = Font(name='Arial', size=10, bold=True)
    c.fill = PatternFill('solid', fgColor=CHIARO)

# ── 2. Articoli ─────────────────────────────────────────────────
wa = wb.create_sheet('Articoli')
intesta(wa, [('Data', 12), ('Cliente', 26), ('Linea', 20), ('Prodotto', 46), ('Quantita', 10)])
r = 2
for v in voci:
    for a in (v.get('articoli') or []):
        scrivi(wa, r, [data_breve(v.get('ts', '')), v.get('nome', ''),
                       LINEE.get(v.get('linea', ''), v.get('linea', '')),
                       a.get('nome', ''), int(a.get('quantita') or 0) or None])
        r += 1
if r > 2:
    wa.cell(row=r + 1, column=4, value='Totale pezzi').font = Font(name='Arial', size=10, bold=True)
    c = wa.cell(row=r + 1, column=5, value='=SUM(E2:E%d)' % (r - 1))
    c.font = Font(name='Arial', size=10, bold=True)
    c.fill = PatternFill('solid', fgColor=CHIARO)

# ── 3. Da programmare ───────────────────────────────────────────
wp = wb.create_sheet('Da programmare')
intesta(wp, [('Prossima', 12), ('Cliente', 26), ('Azienda / Ente', 26), ('Linea', 20),
             ('Cadenza', 30), ('Email', 30), ('Ultimi articoli', 42)])
ricorrenti = [v for v in voci if (v.get('fornitura') or {}).get('tipo') in ('ricorrente', 'stagionale')]
ricorrenti.sort(key=lambda v: v.get('prossima') or '9999')
for r, v in enumerate(ricorrenti, start=2):
    scrivi(wp, r, [
        v.get('prossima', '') or 'da definire',
        v.get('nome', ''), v.get('azienda', ''),
        LINEE.get(v.get('linea', ''), v.get('linea', '')),
        (v.get('fornitura') or {}).get('descrizione', '').split('\n')[0],
        v.get('email', ''),
        ' · '.join(a.get('nome', '') for a in (v.get('articoli') or [])),
    ])

# ── 4. Clienti ──────────────────────────────────────────────────
wc = wb.create_sheet('Clienti')
intesta(wc, [('Cliente', 26), ('Azienda / Ente', 26), ('Email', 30), ('Richieste', 11),
             ('Pezzi totali', 12), ('Linee', 34), ('Ultima richiesta', 15)])
agg = defaultdict(lambda: {'nome': '', 'azienda': '', 'n': 0, 'pezzi': 0, 'linee': set(), 'ultima': ''})
for v in voci:
    k = (v.get('email') or v.get('nome') or '').lower()
    a = agg[k]
    a['nome'] = v.get('nome', '') or a['nome']
    a['azienda'] = v.get('azienda', '') or a['azienda']
    a['n'] += 1
    a['pezzi'] += sum(int(x.get('quantita') or 0) for x in (v.get('articoli') or []))
    if v.get('linea'):
        a['linee'].add(LINEE.get(v['linea'], v['linea']))
    d = data_breve(v.get('ts', ''))
    a['ultima'] = max(a['ultima'], v.get('ts', ''))
for r, (k, a) in enumerate(sorted(agg.items(), key=lambda x: -x[1]['pezzi']), start=2):
    scrivi(wc, r, [a['nome'], a['azienda'], k, a['n'], a['pezzi'] or None,
                   ' · '.join(sorted(a['linee'])), data_breve(a['ultima'])])

wb.save(USCITA)
print('Richieste esportate : %d' % len(voci))
print('Da programmare      : %d' % len(ricorrenti))
print('Clienti distinti    : %d' % len(agg))
print('Scritto             : %s' % USCITA)
